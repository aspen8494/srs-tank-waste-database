#!/usr/bin/env python3
"""Objective QA battery for the white paper (per skill reference layout-and-render.md)."""
import re, zipfile, sys
import pypdf
import pdfplumber

PDF = "/Users/aspen_mb/Documents/Savannah_RiverTF/paper/SRS_Hanford_Tank_Waste_Comparison_White_Paper.pdf"
DOCX = "/Users/aspen_mb/Documents/Savannah_RiverTF/paper/SRS_Hanford_Tank_Waste_Comparison_White_Paper.docx"

fails = []

# 1. pypdf: page count + headline numbers + LaTeX-token grep + figure sequence
rd = pypdf.PdfReader(PDF)
pages = [p.extract_text() for p in rd.pages]
full = "\n".join(pages)
print(f"PDF pages: {len(pages)}")
if not (10 <= len(pages) <= 20):
    fails.append(f"page count {len(pages)} outside 10-20")

headlines = [
    "33.2 Mgal", "191 MCi", "53.0 Mgal", "122.5 MCi", "5.8 MCi/Mgal",
    "Executive Summary", "1. Introduction", "2. Results", "3. Methodology",
    "4. Discussion", "5. Conclusions", "References",
    "2.5-fold", "156 percent", "5.80 MCi", "2.70 MCi", "299 t", "150,200",
    "Sm-151", "Zr-93", "Pu-238", "Tank 16", "Tanks 18 and 19", "Salt Batch 14", "I-129",
]
for h in headlines:
    if h not in full:
        fails.append(f"headline missing from PDF text: {h!r}")

for tok in (r"\sum", r"\lceil", r"\rceil", r"\frac"):
    if tok in full:
        fails.append(f"LaTeX token leaked: {tok}")

figseq = re.findall(r"Figure (\d+)\.", full)
firsts = []
for n in figseq:
    if n not in firsts:
        firsts.append(n)
if firsts != [str(i) for i in range(1, 6)]:
    fails.append(f"figure first-occurrence sequence wrong: {firsts}")

# 2. pdfplumber: every table row column present (spot-check key rows)
with pdfplumber.open(PDF) as pdf:
    txt = "\n".join((p.extract_text() or "") for p in pdf.pages)
for probe in ("33.2", "122.5", "5.8", "2.3", "0.63", "1.56",
              "1.46E+06", "2.94E+06", "19.3", "566", "16,600", "5,870", "1.15E+05"):
    if probe not in txt:
        fails.append(f"table cell missing from PDF text layer: {probe!r}")

# 3. margins: render a mid-page and measure bbox (pypdfium2 if present)
try:
    import pypdfium2 as pdfium
    import numpy as np
    from PIL import Image as PILImage
    pg = pdfium.PdfDocument(PDF)[len(pages) // 2]
    bmp = pg.render(scale=1.2)
    arr = np.asarray(bmp.to_pil().convert("L"))
    mask = arr < 200
    cols = np.where(mask.any(axis=0))[0]
    rows_ = np.where(mask.any(axis=1))[0]
    left_px = cols[0]; right_px = arr.shape[1] - cols[-1]
    print(f"mid-page margins: left {left_px}px right {right_px}px (page {len(pages)//2+1})")
    if abs(left_px - right_px) > 15:
        fails.append(f"asymmetric margins: left {left_px} right {right_px}")
except ImportError:
    print("pypdfium2 unavailable; margin check skipped")

# 4. docx readback: Normal style, counts, zip integrity
import docx as pydocx
d = pydocx.Document(DOCX)
normal = d.styles["Normal"]
assert normal.font.name == "Times New Roman", f"Normal font is {normal.font.name}"
assert normal.font.size.pt == 12, f"Normal size is {normal.font.size.pt}"
h1 = [p for p in d.paragraphs if p.style.name == "Heading 1"]
imgs = len(d.inline_shapes)
tbls = len(d.tables)
print(f"docx: Normal=Times 12pt | headings={len(h1)} | images={imgs} | tables={tbls}")
if imgs != 5:
    fails.append(f"docx images {imgs} != 5")
if tbls != 4:
    fails.append(f"docx tables {tbls} != 4")
z = zipfile.ZipFile(DOCX)
if z.testzip() is not None:
    fails.append("docx zip integrity failed")

# 5. caption adjacency: every image followed by its Figure N line on the same page (pymupdf)
try:
    import fitz
    fdoc = fitz.open(PDF)
    for pno in range(len(fdoc)):
        pg = fdoc[pno]
        blocks = pg.get_text("dict")["blocks"]
        img_y = []
        cap_y = []
        for b in blocks:
            if b["type"] == 1:
                img_y.append(b["bbox"][1])
        lines = []
        for b in blocks:
            if b["type"] == 0:
                for ln in b["lines"]:
                    s = "".join(sp["text"] for sp in ln["spans"])
                    if s.strip().startswith("Figure"):
                        lines.append((ln["bbox"][1], s.strip()[:30]))
        # each image must have a caption below it on the same page
        for iy in img_y:
            if not any(cy > iy for cy, _ in lines):
                fails.append(f"page {pno+1}: image without caption below on same page")
    print("caption adjacency checked")
except ImportError:
    print("pymupdf unavailable; adjacency check skipped")

# verdict
print("\n=== QA RESULT:", "PASS" if not fails else f"{len(fails)} FAILURES", "===")
for f_ in fails:
    print(" -", f_)
sys.exit(1 if fails else 0)
